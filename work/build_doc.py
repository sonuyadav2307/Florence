from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path('/Users/sonu/Documents/Codex/2026-09-20/i-w-2')
source = (ROOT / 'outputs/Florence_Build_Specification.md').read_text()
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(.70)
section.bottom_margin = Inches(.70)
section.left_margin = Inches(.72)
section.right_margin = Inches(.72)
section.footer_distance = Inches(.30)
width = 6.83

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.widow_control = True
    st.paragraph_format.space_after = Pt(6)
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.12
doc.styles['Title'].font.size = Pt(28)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(12)
doc.styles['Subtitle'].font.size = Pt(10)
for name, size in [('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(14)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
for name in ['List Bullet', 'List Number']:
    doc.styles[name].paragraph_format.space_after = Pt(4)
    doc.styles[name].paragraph_format.line_spacing = 1.08

def hyperlink(p, label, url):
    rel = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), rel)
    r = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    col = OxmlElement('w:color'); col.set(qn('w:val'), '24553C')
    rp.append(col)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rp.append(u)
    r.append(rp)
    t = OxmlElement('w:t'); t.text = label; r.append(t)
    h.append(r); p._p.append(h)

def inline(p, s):
    pat = r'(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for token in re.split(pat, s):
        if not token: continue
        if token.startswith('[') and '](' in token:
            m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
            hyperlink(p, m[1], m[2])
        elif token.startswith('**'):
            p.add_run(token[2:-2]).bold = True
        elif token.startswith('`'):
            run = p.add_run(token[1:-1]); run.font.name='Liberation Mono'; run.font.size=Pt(9)
        else: p.add_run(token)

def make_table(rows):
    n = len(rows[0])
    t = doc.add_table(rows=0, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    headers = rows[0]
    if headers[0] == 'Mode': weights=[1.35,.92,.92,1.80,1.80]
    elif headers[0] == 'ID' and n == 6: weights=[1.15,1.35,.77,1.15,.85,1.1]
    elif headers[0] == 'ID': weights=[.70,2.35,3.78]
    elif headers[0] == 'Table': weights=[1.3,5.53]
    elif headers[0] == 'Route': weights=[2.5,2.43,1.9]
    elif headers[0] == 'Method and route': weights=[2.25,2.2,2.38]
    elif headers[0] == 'Preset': weights=[1.35,1.096,1.096,1.096,1.096,1.096]
    elif n == 3: weights=[1.2,1.8,3.83]
    elif n == 4: weights=[1.2,1.9,2,1.73]
    else: weights=[1]*n
    ws=[width*x/sum(weights) for x in weights]
    for col,w in zip(t.columns,ws): col.width=Inches(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'D9D9D9'); borders.append(edge)
    props.append(borders)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        trp=t.rows[-1]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit'); trp.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); trp.append(repeat)
        for ci,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Inches(ws[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'333B37' if ri==0 else ('F5F6F5' if ri%2==0 else 'FFFFFF')); tc.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'85'); el.set(qn('w:type'),'dxa'); margins.append(el)
            tc.append(margins)
            p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1); p.paragraph_format.line_spacing=1.04
            inline(p,value)
            for run in p.runs:
                run.font.size=Pt(9)
                if ri==0: run.font.bold=True; run.font.color.rgb=RGBColor(255,255,255)
            if ri==0: p.paragraph_format.keep_with_next=True
            if headers[0]=='Preset' and ci>0:
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=0.3

lines=source.splitlines(); i=0
while i<len(lines):
    line=lines[i]
    if not line.strip(): i+=1; continue
    if line.startswith('```'):
        block=[]; i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            block.append(lines[i]); i+=1
        for idx,text in enumerate(block):
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0
            p.paragraph_format.keep_with_next=idx<len(block)-1 and len(block)<18
            r=p.add_run(text or ' '); r.font.name='Liberation Mono'; r.font.size=Pt(8.4)
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=.3
        i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',x) for x in cells): rows.append(cells)
            i+=1
        make_table(rows); continue
    m=re.match(r'^(#{1,4}) (.*)',line)
    if m:
        level=len(m[1]); label=m[2]
        style='Title' if level==1 else 'Heading '+str(level-1)
        p=doc.add_paragraph(style=style); inline(p,label)
    elif line.startswith('Version 1.0'):
        p=doc.add_paragraph(style='Subtitle'); inline(p,line)
    elif line.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); inline(p,line[2:])
    elif re.match(r'^\d+\. ',line):
        # Preserve specified ordinal text, avoiding Word's cross-list continuation.
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.20); p.paragraph_format.first_line_indent=Inches(-.20)
        inline(p,line)
    else:
        p=doc.add_paragraph(); inline(p,line)
    i+=1

footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
rr=footer.add_run('Florence Build Specification  |  '); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(0,0,0)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title='Florence Web App Build Specification'
doc.core_properties.subject='Product requirements and implementation contract'
doc.core_properties.author=''
for root in [doc.styles.element, doc.element]:
    for node in list(root.iter(qn('w:pBdr'))):
        node.getparent().remove(node)
doc.save(ROOT/'outputs/Florence_Build_Specification.docx')
print('DOCX created')
